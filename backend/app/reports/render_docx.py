from __future__ import annotations

import base64
import io
import re
from pathlib import Path
from typing import Any

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor

from .models import (
    CalloutBlock,
    ChartBlock,
    GraphImageBlock,
    KPIsBlock,
    ReportModel,
    TableBlock,
    TextBlock,
)

# Cores institucionais do COGNIS para DOCX
COLOR_PRIMARY = RGBColor(0x4F, 0x46, 0xE5)  # Índigo / Accent
COLOR_TEXT_MAIN = RGBColor(0x0F, 0x17, 0x2A)  # Ardósia escuro
COLOR_TEXT_MUTED = RGBColor(0x64, 0x74, 0x8B)  # Texto secundário
HEX_HEADER_BG = "F1F5F9"  # Cinza suave para cabeçalho de tabela
HEX_CALLOUT_INFO_BG = "F0F9FF"
HEX_CALLOUT_WARN_BG = "FFFBEB"


def _set_cell_background(cell: Any, fill_hex: str) -> None:
    """Aplica cor de fundo em célula de tabela via OXML de forma compatível com Word."""
    try:
        shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
        cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))
    except Exception:
        pass


def _resolve_image_stream(png_path: str | None) -> io.BytesIO | str | None:
    """
    Resolve um caminho de imagem ou data URI para um stream de bytes ou caminho de arquivo legível.
    """
    if not png_path:
        return None

    if png_path.startswith("data:image/"):
        try:
            _, b64_data = png_path.split(",", 1)
            raw = base64.b64decode(b64_data)
            return io.BytesIO(raw)
        except Exception:
            return None

    p = Path(png_path)
    if p.is_file():
        return str(p)

    return None


def render_docx_report(
    report: ReportModel | dict[str, Any],
    output_path: str | Path | None = None,
) -> bytes:
    """
    Renderiza um relatório COGNIS (ReportModel) em documento DOCX formatado (SPEC §10.3).
    - Títulos, metadados e banner de conformidade ética (§13.4).
    - Tabelas de KPIs e de dados formatadas com estilo nativo limpo ('Table Grid').
    - Imagens de grafo e gráficos embutidos em resolução proporcional com legendas.
    - Textos enriquecidos com selos de proveniência (IA Local, Regras, Dados).
    - Apêndice com definições de fórmulas e limitações de dados.
    - Se output_path for informado, grava no disco em formato .docx.
    - Retorna os bytes do arquivo DOCX.
    """
    if isinstance(report, dict):
        report_model = ReportModel.model_validate(report)
    else:
        report_model = report

    doc = docx.Document()

    # Configuração de margens (1 polegada / 2.54 cm)
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)

    # 1. Cabeçalho e Título
    brand_para = doc.add_paragraph()
    brand_run = brand_para.add_run("COGNIS — INTELIGÊNCIA DE COLABORAÇÃO")
    brand_run.font.size = Pt(9)
    brand_run.font.bold = True
    brand_run.font.color.rgb = COLOR_PRIMARY

    title_para = doc.add_paragraph()
    title_run = title_para.add_run(report_model.meta.title)
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = COLOR_TEXT_MAIN
    title_para.paragraph_format.space_after = Pt(8)

    # Metadados em tabela de 2 colunas
    meta_table = doc.add_table(rows=0, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    meta_items = [
        ("Tipo de Relatório", f"{report_model.meta.type.capitalize()}{' (Anonimizado)' if report_model.meta.anonymized else ''}"),
        ("Projeto", report_model.meta.project),
        (
            "Período",
            f"{report_model.meta.window.get('label') or ''} ({report_model.meta.window.get('from', '—')} até {report_model.meta.window.get('to', '—')})".strip(),
        ),
        ("Gerado em", report_model.meta.generated_at),
        ("Versão dos Dados", str(report_model.meta.data_version or "1.0")),
    ]
    if report_model.meta.person_name:
        meta_items.append(("Colaborador", report_model.meta.person_name))

    for label, val in meta_items:
        row = meta_table.add_row()
        lbl_cell = row.cells[0]
        val_cell = row.cells[1]
        lbl_p = lbl_cell.paragraphs[0]
        val_p = val_cell.paragraphs[0]

        lbl_run = lbl_p.add_run(label)
        lbl_run.font.bold = True
        lbl_run.font.size = Pt(9.5)
        lbl_run.font.color.rgb = COLOR_TEXT_MUTED

        val_run = val_p.add_run(val)
        val_run.font.size = Pt(9.5)
        val_run.font.color.rgb = COLOR_TEXT_MAIN

        _set_cell_background(lbl_cell, "F8FAFC")
        _set_cell_background(val_cell, "F8FAFC")

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Banner Ético / Uso Responsável (§13.4)
    ethical_p = doc.add_paragraph()
    ethical_p.paragraph_format.left_indent = Inches(0.25)
    ethical_p.paragraph_format.space_after = Pt(14)
    e_tag = ethical_p.add_run("Nota de Uso Responsável: ")
    e_tag.font.bold = True
    e_tag.font.size = Pt(9.5)
    e_tag.font.color.rgb = COLOR_PRIMARY
    e_text = ethical_p.add_run(report_model.appendix.ethical_note)
    e_text.font.italic = True
    e_text.font.size = Pt(9.5)
    e_text.font.color.rgb = COLOR_TEXT_MUTED

    # 2. Seções do Relatório
    for section in report_model.sections:
        h1 = doc.add_heading(level=1)
        h1_run = h1.add_run(section.title)
        h1_run.font.size = Pt(15)
        h1_run.font.bold = True
        h1_run.font.color.rgb = COLOR_PRIMARY
        h1.paragraph_format.space_before = Pt(14)
        h1.paragraph_format.space_after = Pt(6)

        for block in section.blocks:
            if isinstance(block, KPIsBlock):
                items = block.items
                if items:
                    kpi_table = doc.add_table(rows=1, cols=3, style="Table Grid")
                    kpi_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    hdr_cells = kpi_table.rows[0].cells
                    for idx, h_name in enumerate(["Indicador", "Valor", "Referência / Meta"]):
                        c = hdr_cells[idx]
                        _set_cell_background(c, HEX_HEADER_BG)
                        p = c.paragraphs[0]
                        r = p.add_run(h_name)
                        r.font.bold = True
                        r.font.size = Pt(9.5)

                    for item in items:
                        r_cells = kpi_table.add_row().cells
                        # Indicador
                        lbl_p = r_cells[0].paragraphs[0]
                        lbl_r = lbl_p.add_run(item.label)
                        lbl_r.font.size = Pt(9.5)

                        # Valor
                        val_p = r_cells[1].paragraphs[0]
                        val_r = val_p.add_run(str(item.value))
                        val_r.font.bold = True
                        val_r.font.size = Pt(11)
                        val_r.font.color.rgb = COLOR_PRIMARY

                        # Hint
                        hint_p = r_cells[2].paragraphs[0]
                        hint_r = hint_p.add_run(item.hint or "—")
                        hint_r.font.size = Pt(9)
                        hint_r.font.color.rgb = COLOR_TEXT_MUTED

                    doc.add_paragraph().paragraph_format.space_after = Pt(6)

            elif isinstance(block, TableBlock):
                if block.title:
                    tbl_title_p = doc.add_paragraph()
                    tt_run = tbl_title_p.add_run(block.title)
                    tt_run.font.bold = True
                    tt_run.font.size = Pt(11)
                    tbl_title_p.paragraph_format.space_after = Pt(3)

                cols_count = len(block.columns)
                if cols_count > 0:
                    data_table = doc.add_table(rows=1, cols=cols_count, style="Table Grid")
                    data_table.alignment = WD_TABLE_ALIGNMENT.CENTER

                    hdr_cells = data_table.rows[0].cells
                    for c_idx, col_name in enumerate(block.columns):
                        c = hdr_cells[c_idx]
                        _set_cell_background(c, HEX_HEADER_BG)
                        p = c.paragraphs[0]
                        r = p.add_run(col_name)
                        r.font.bold = True
                        r.font.size = Pt(9.5)

                    for row_data in block.rows:
                        row_cells = data_table.add_row().cells
                        for c_idx, val in enumerate(row_data):
                            if c_idx < len(row_cells):
                                p = row_cells[c_idx].paragraphs[0]
                                r = p.add_run(str(val))
                                r.font.size = Pt(9.5)

                    doc.add_paragraph().paragraph_format.space_after = Pt(6)

            elif isinstance(block, (ChartBlock, GraphImageBlock)):
                img_stream = _resolve_image_stream(block.png_path)
                if img_stream is not None:
                    pic_p = doc.add_paragraph()
                    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    try:
                        pic_p.add_run().add_picture(img_stream, width=Inches(5.8))
                    except Exception:
                        pic_p.add_run(f"[Erro ao carregar imagem: {block.caption}]").font.italic = True

                    if block.caption:
                        cap_p = doc.add_paragraph()
                        cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        cap_p.paragraph_format.space_after = Pt(8)
                        cap_r = cap_p.add_run(block.caption)
                        cap_r.font.italic = True
                        cap_r.font.size = Pt(9)
                        cap_r.font.color.rgb = COLOR_TEXT_MUTED
                else:
                    placeholder_p = doc.add_paragraph()
                    placeholder_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    pr = placeholder_p.add_run(f"[ Imagem não disponível: {block.caption or 'gráfico'} ]")
                    pr.font.italic = True
                    pr.font.size = Pt(9.5)
                    pr.font.color.rgb = COLOR_TEXT_MUTED

            elif isinstance(block, TextBlock):
                origin_p = doc.add_paragraph()
                origin_p.paragraph_format.space_before = Pt(4)
                origin_p.paragraph_format.space_after = Pt(2)
                badge_text = (
                    "✦ ANÁLISE IA LOCAL (Llama 3.1 8B)"
                    if block.origin == "llm"
                    else ("⚙ ANÁLISE DETERMINÍSTICA (Regras)" if block.origin == "template" else "FATOS E DADOS")
                )
                badge_run = origin_p.add_run(badge_text)
                badge_run.font.size = Pt(8.5)
                badge_run.font.bold = True
                badge_run.font.color.rgb = COLOR_PRIMARY if block.origin != "llm" else RGBColor(0x05, 0x96, 0x69)

                # Processamento de linhas com suporte a bullets e negrito básico
                raw_lines = block.markdown.split("\n")
                for line in raw_lines:
                    s_line = line.strip()
                    if not s_line:
                        continue
                    if s_line.startswith("- "):
                        p = doc.add_paragraph(style="List Bullet")
                        content = s_line[2:].strip()
                    else:
                        p = doc.add_paragraph()
                        content = s_line

                    # Divisão básica de negrito **texto**
                    parts = re.split(r"(\*\*.*?\*\*)", content)
                    for part in parts:
                        if part.startswith("**") and part.endswith("**"):
                            r = p.add_run(part[2:-2])
                            r.font.bold = True
                        else:
                            r = p.add_run(part)
                        r.font.size = Pt(10)
                        r.font.color.rgb = COLOR_TEXT_MAIN

            elif isinstance(block, CalloutBlock):
                callout_p = doc.add_paragraph()
                callout_p.paragraph_format.left_indent = Inches(0.2)
                callout_p.paragraph_format.space_before = Pt(4)
                callout_p.paragraph_format.space_after = Pt(6)

                prefix_map = {
                    "warning": "[AVISO DE ATENÇÃO] ",
                    "success": "[DESTAQUE POSITIVO] ",
                    "info": "[INFORMAÇÃO] ",
                }
                c_prefix = prefix_map.get(block.level, "[INFO] ")
                tag_r = callout_p.add_run(c_prefix)
                tag_r.font.bold = True
                tag_r.font.size = Pt(9.5)
                tag_r.font.color.rgb = (
                    RGBColor(0xD9, 0x77, 0x06)
                    if block.level == "warning"
                    else (RGBColor(0x16, 0xA3, 0x4A) if block.level == "success" else COLOR_PRIMARY)
                )

                txt_r = callout_p.add_run(block.text)
                txt_r.font.size = Pt(9.5)
                txt_r.font.color.rgb = COLOR_TEXT_MAIN

    # 3. Apêndice e Metodologia
    doc.add_page_break()
    app_h = doc.add_heading(level=1)
    app_h_run = app_h.add_run("Apêndice e Metodologia")
    app_h_run.font.size = Pt(15)
    app_h_run.font.bold = True
    app_h_run.font.color.rgb = COLOR_PRIMARY

    if report_model.appendix.definitions:
        sub_h = doc.add_heading(level=2)
        sub_h_r = sub_h.add_run("Definições das Métricas")
        sub_h_r.font.size = Pt(12)
        sub_h_r.font.bold = True
        sub_h_r.font.color.rgb = COLOR_TEXT_MAIN

        for def_str in report_model.appendix.definitions:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(def_str)
            r.font.size = Pt(9)
            r.font.color.rgb = COLOR_TEXT_MUTED

    if report_model.appendix.limitations:
        sub_h2 = doc.add_heading(level=2)
        sub_h2_r = sub_h2.add_run("Limitações dos Dados")
        sub_h2_r.font.size = Pt(12)
        sub_h2_r.font.bold = True
        sub_h2_r.font.color.rgb = COLOR_TEXT_MAIN

        for lim_str in report_model.appendix.limitations:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(lim_str)
            r.font.size = Pt(9)
            r.font.color.rgb = COLOR_TEXT_MUTED

    # Rodapé final
    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.paragraph_format.space_before = Pt(20)
    fr = footer_p.add_run(
        "COGNIS — Análise determinística e interpretação contextual de grafos de colaboração."
    )
    fr.font.size = Pt(8.5)
    fr.font.italic = True
    fr.font.color.rgb = COLOR_TEXT_MUTED

    # Salvar em buffer de bytes
    bio = io.BytesIO()
    doc.save(bio)
    docx_bytes = bio.getvalue()

    if output_path is not None:
        p = Path(output_path)
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(docx_bytes)

    return docx_bytes
